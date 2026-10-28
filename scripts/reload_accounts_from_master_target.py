#!/usr/bin/env python3
"""
Build CSVs from Organizations Final Master List + Target List (All Team),
then bulk-upsert Accounts, Affiliations, and Account Territory Fields.

Usage:
  python3 ph/scripts/reload_accounts_from_master_target.py prepare
  python3 ph/scripts/reload_accounts_from_master_target.py load-accounts
  python3 ph/scripts/reload_accounts_from_master_target.py load-affiliations
  python3 ph/scripts/reload_accounts_from_master_target.py load-targets
  python3 ph/scripts/reload_accounts_from_master_target.py report
  python3 ph/scripts/reload_accounts_from_master_target.py all   # prepare + loads after wipe is at 0
"""

from __future__ import annotations

import csv
import json
import re
import subprocess
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "ph" / "scripts" / "tmp" / "account-reload"
MASTER_XLSX = Path(
    "/Users/ashrafrezk/Downloads/Target_&_Master_list/Organizations_FinalMaster_List_flagged.xlsx"
)
TARGET_XLSX = Path(
    "/Users/ashrafrezk/Downloads/Target_&_Master_list/Target List (All Team).xlsx"
)

# Queried from org (PersonAccount IsPersonType=true preferred).
RT = {
    "Institution_HCO": "012hk000000Ed37AAC",
    "Pharmacy": "012hk000000Ed39AAC",
    "Distributor": "012hk000000EdRLAA0",
    "PersonAccount": "012hk000000Ed4jAAC",
}

LATIN_WORKPLACE = {
    "ABB",
    "GUC",
    "PGESCO",
    "CPC",
    "PAS",
    "P.M.S",
    "PMS",
    "E GAS",
    "EGAS",
    "OGS",
    "IDA",
    "DMS",
    "DUSIT",
    "THARWA",
    "THAWRA",
    "UCP",
    "AUC",
    "AXA",
    "DOMTY",
    "GASTEC",
    "HYPER",
    "MOPCO",
    "PETROMINT",
    "PETROSENNAN",
    "THERMOTECH",
    "VACSERA",
    "ZETCO",
    "SAN MASR",
    "PRIMA",
}
WORKPLACE_PREFIXES = [
    "قسم",
    "بنك",
    "شركه",
    "شركة",
    "مصنع",
    "كهرباء",
    "كهربا",
    "تعاقد",
    "هيئه",
    "هيئة",
    "وزاره",
    "وزارة",
    "جهاز",
    "مستشفى",
    "مستشفي",
    "هايبر",
    "مرفق",
]
ARABIC_RE = re.compile(r"[\u0600-\u06FF]")
LATIN_NAME_RE = re.compile(r"^[A-Za-z0-9 .,&/+\-']+$")
LEADING_LATIN_RE = re.compile(r"^([A-Za-z][A-Za-z0-9.&]{1,})\b")


def norm(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def doctor_ext(code: str) -> str:
    return "DOC-" + norm(code)


def clinic_ext(clinic_id: str) -> str:
    return "CLINIC-" + norm(clinic_id)


def aff_ext(doctor_code: str, related_ext: str) -> str:
    return "AFF-" + doctor_ext(doctor_code) + "-" + norm(related_ext)


def map_tier(customer_class: str) -> str | None:
    upper = norm(customer_class).upper()
    for ch in upper:
        if ch in ("A", "B", "C"):
            return ch
    return None


def to_bool(raw, default=True) -> bool:
    if raw is None or raw == "":
        return default
    if isinstance(raw, bool):
        return raw
    s = str(raw).strip().lower()
    if s in ("true", "1", "yes", "y"):
        return True
    if s in ("false", "0", "no", "n"):
        return False
    return default


def normalize_phone(raw) -> str | None:
    if raw is None or raw == "":
        return None
    digits = re.sub(r"[^0-9+]", "", str(raw))
    if not digits:
        return None
    if not digits.startswith("+") and len(digits) == 10 and digits.startswith("1"):
        return "+20" + digits
    if not digits.startswith("+") and len(digits) >= 10:
        return "+20" + digits
    return digits[:40]


def clean_prefixed_name(raw: str) -> str:
    if not raw or "_" not in raw:
        return raw
    idx = raw.index("_")
    prefix, rest = raw[:idx], raw[idx + 1 :].strip()
    if prefix.isdigit() and rest:
        return rest[:255]
    return raw


def clinic_display_name(doctor_name: str, area: str, clinic_id: str) -> str:
    doc = norm(doctor_name) or ("Clinic " + clinic_id)
    loc = norm(area) or "غير محدد"
    return ("عيادة " + doc + " - " + loc)[:255]


def split_person_name(full_name: str) -> tuple[str | None, str]:
    cleaned = re.sub(r"\s+", " ", norm(full_name))
    if not cleaned:
        return None, "Unknown"
    space = cleaned.find(" ")
    if space <= 0:
        return None, cleaned[:80]
    return cleaned[:space][:40], cleaned[space + 1 :][:80]


def is_workplace_doctor_name(doctor_name: str) -> bool:
    name = norm(doctor_name)
    if not name:
        return False
    upper = name.upper()
    compact = upper.replace(" ", "").replace(".", "")
    if upper in LATIN_WORKPLACE or compact in LATIN_WORKPLACE:
        return True
    if not ARABIC_RE.search(name) and LATIN_NAME_RE.match(name):
        return True
    m = LEADING_LATIN_RE.match(name)
    if m:
        token = m.group(1)
        rest = name[m.end() :].strip()
        up = token.upper().replace(".", "")
        if token.upper() in LATIN_WORKPLACE or up in LATIN_WORKPLACE:
            return True
        if rest and ARABIC_RE.search(rest) and len(token) >= 2:
            return True
    for prefix in WORKPLACE_PREFIXES:
        if name.startswith(prefix):
            return True
    return False


def read_sheet1(path: Path) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    headers = [norm(h) if h is not None else "" for h in next(rows_iter)]
    out = []
    for values in rows_iter:
        rec = {}
        empty = True
        for i, h in enumerate(headers):
            if not h:
                continue
            v = values[i] if i < len(values) else None
            if v is not None and str(v).strip() != "":
                empty = False
            rec[h] = v
        if not empty:
            out.append(rec)
    wb.close()
    return out


def write_csv(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        # Bulk API lineEnding=LF requires Unix newlines (Python csv defaults to \r\n).
        w = csv.DictWriter(
            f, fieldnames=fieldnames, extrasaction="ignore", lineterminator="\n"
        )
        w.writeheader()
        for row in rows:
            w.writerow({k: ("" if row.get(k) is None else row.get(k)) for k in fieldnames})


def sf_json(args: list[str]) -> dict:
    proc = subprocess.run(
        ["sf", *args, "--json"],
        cwd=str(ROOT),
        capture_output=True,
        text=True,
    )
    try:
        data = json.loads(proc.stdout or "{}")
    except json.JSONDecodeError:
        print(proc.stdout)
        print(proc.stderr)
        raise
    if proc.returncode != 0 and data.get("status") not in (0, None):
        raise RuntimeError(json.dumps(data, indent=2)[:4000])
    return data


def wait_bulk_job(job_id: str, timeout_s: int = 7200) -> dict:
    start = time.time()
    while time.time() - start < timeout_s:
        data = sf_json(["data", "bulk", "results", "--job-id", job_id])
        result = data.get("result") or data
        # CLI shapes vary; also poll via query on job.
        status_data = sf_json(
            [
                "data",
                "query",
                "--query",
                f"SELECT Id, State, NumberRecordsProcessed, NumberRecordsFailed FROM JobInfo WHERE Id='{job_id}'",
                "--use-tooling-api",
            ]
        )
        records = (status_data.get("result") or {}).get("records") or []
        if records:
            state = records[0].get("State")
            print(
                f"  job {job_id} state={state} processed={records[0].get('NumberRecordsProcessed')} "
                f"failed={records[0].get('NumberRecordsFailed')}"
            )
            if state in ("JobComplete", "Failed", "Aborted"):
                return records[0]
        else:
            # Fallback: ingest job object via data bulk status if available
            print(f"  waiting on {job_id}…")
        time.sleep(15)
    raise TimeoutError(f"Bulk job {job_id} timed out")


def upsert_csv(sobject: str, external_id: str, csv_path: Path, wait: bool = True) -> str:
    print(f"Upserting {csv_path.name} → {sobject} on {external_id}")
    # --wait in minutes; large files need a long wait.
    wait_min = "180" if wait else "0"
    data = sf_json(
        [
            "data",
            "upsert",
            "bulk",
            "--sobject",
            sobject,
            "--external-id",
            external_id,
            "--file",
            str(csv_path),
            "--line-ending",
            "LF",
            "--wait",
            wait_min,
        ]
    )
    result = data.get("result") or {}
    job_id = result.get("jobId") or result.get("id")
    if not job_id:
        job_id = (result.get("jobs") or [{}])[0].get("jobId")
    # When --wait completes, jobId may still be present; also accept success without id.
    state = result.get("state") or result.get("Status")
    failed = result.get("numberRecordsFailed")
    processed = result.get("numberRecordsProcessed")
    print(f"  job={job_id} state={state} processed={processed} failed={failed}")
    if failed and int(failed) > 0:
        print("  WARNING: some rows failed — check sf data upsert resume / job results")
    if not job_id and data.get("status") not in (0, None):
        raise RuntimeError("Upsert failed: " + json.dumps(data)[:2000])
    return job_id or ""


def prepare() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    print("Reading master…")
    master = read_sheet1(MASTER_XLSX)
    print(f"  {len(master)} rows")

    org_rows: list[dict] = []
    person_rows: list[dict] = []
    aff_rows: list[dict] = []
    doctors_by_code: dict[str, dict] = {}
    workplace_skipped: list[dict] = []
    clinic_ids: set[str] = set()
    doc_codes_created: set[str] = set()

    for rec in master:
        type_name = norm(rec.get("Organization Type Name"))
        code = norm(rec.get("Code"))
        doctor_code = norm(rec.get("Doctor Code"))
        clinic_id = norm(rec.get("Doctor Clinic ID"))
        name = norm(rec.get("Name"))
        area = norm(rec.get("Area"))
        address = norm(rec.get("Address"))
        customer_class = norm(rec.get("Class"))
        active = to_bool(rec.get("Active"), True)
        status = "Active" if active else "Inactive"
        phone = normalize_phone(rec.get("Telephone"))
        email = norm(rec.get("Email")) or None
        tier = map_tier(customer_class)

        if type_name.lower() == "doctor":
            if not clinic_id:
                continue
            clinic_ids.add(clinic_id)
            org_rows.append(
                {
                    "External_ID__c": clinic_ext(clinic_id),
                    "RecordTypeId": RT["Institution_HCO"],
                    "Name": clinic_display_name(name, area, clinic_id),
                    "Institution_Type__c": "Clinic",
                    "BillingStreet": address or None,
                    "BillingCity": area or None,
                    "BillingCountry": "Egypt",
                    "Customer_Class__c": customer_class or None,
                    "Tier__c": tier,
                    "Account_Status__c": status,
                    "Phone": phone,
                }
            )
            if doctor_code and doctor_code not in doctors_by_code:
                doctors_by_code[doctor_code] = {
                    "code": doctor_code,
                    "doctor": name,
                    "customerClass": customer_class,
                    "active": active,
                    "mobile": phone,
                    "email": email,
                }
            if doctor_code and clinic_id:
                aff_rows.append(
                    {
                        "External_ID__c": aff_ext(doctor_code, clinic_ext(clinic_id)),
                        "Primary_Account__r.External_ID__c": doctor_ext(doctor_code),
                        "Related_Account__r.External_ID__c": clinic_ext(clinic_id),
                        "Affiliation_Type__c": "Practice",
                        "Role__c": "Physician",
                        "Is_Active__c": "true",
                    }
                )
            continue

        if type_name.lower() == "pharmacy":
            clean_name = clean_prefixed_name(name)
            if not code or not clean_name:
                continue
            org_rows.append(
                {
                    "External_ID__c": code,
                    "RecordTypeId": RT["Pharmacy"],
                    "Name": clean_name[:255],
                    "Institution_Type__c": "",
                    "BillingStreet": address or None,
                    "BillingCity": area or None,
                    "BillingCountry": "Egypt",
                    "Customer_Class__c": customer_class or None,
                    "Tier__c": tier,
                    "Account_Status__c": status,
                    "Phone": phone,
                }
            )
        elif type_name.lower() == "distributor":
            if not code or not name:
                continue
            org_rows.append(
                {
                    "External_ID__c": code,
                    "RecordTypeId": RT["Distributor"],
                    "Name": name[:255],
                    "Institution_Type__c": "",
                    "BillingStreet": address or None,
                    "BillingCity": area or None,
                    "BillingCountry": "Egypt",
                    "Customer_Class__c": customer_class or None,
                    "Tier__c": tier,
                    "Account_Status__c": status,
                    "Phone": phone,
                }
            )
        elif type_name.lower() in ("hospital", "contract"):
            if not code or not name:
                continue
            org_rows.append(
                {
                    "External_ID__c": code,
                    "RecordTypeId": RT["Institution_HCO"],
                    "Name": name[:255],
                    "Institution_Type__c": "Hospital"
                    if type_name.lower() == "hospital"
                    else "Contract",
                    "BillingStreet": address or None,
                    "BillingCity": area or None,
                    "BillingCountry": "Egypt",
                    "Customer_Class__c": customer_class or None,
                    "Tier__c": tier,
                    "Account_Status__c": status,
                    "Phone": phone,
                }
            )

    for doc in doctors_by_code.values():
        if is_workplace_doctor_name(doc["doctor"]):
            workplace_skipped.append(doc)
            continue
        first, last = split_person_name(doc["doctor"])
        person_rows.append(
            {
                "External_ID__c": doctor_ext(doc["code"]),
                "RecordTypeId": RT["PersonAccount"],
                "FirstName": first,
                "LastName": last,
                "Institution_Type__c": "Doctor",
                "Customer_Class__c": doc["customerClass"] or None,
                "Tier__c": map_tier(doc["customerClass"] or ""),
                "Account_Status__c": "Active" if doc["active"] else "Inactive",
                "BillingCountry": "Egypt",
                "PersonMobilePhone": doc["mobile"],
                "PersonEmail": doc["email"],
            }
        )
        doc_codes_created.add(doc["code"])

    # Drop affiliations whose HCP was skipped as workplace.
    aff_rows = [
        a
        for a in aff_rows
        if norm(a["Primary_Account__r.External_ID__c"]).removeprefix("DOC-") in doc_codes_created
    ]

    # Deduplicate org/person by External_ID__c (last wins).
    def dedupe(rows: list[dict]) -> list[dict]:
        by = {}
        for r in rows:
            by[r["External_ID__c"]] = r
        return list(by.values())

    org_rows = dedupe(org_rows)
    person_rows = dedupe(person_rows)
    aff_by = {a["External_ID__c"]: a for a in aff_rows}
    aff_rows = list(aff_by.values())

    org_fields = [
        "External_ID__c",
        "RecordTypeId",
        "Name",
        "Institution_Type__c",
        "BillingStreet",
        "BillingCity",
        "BillingCountry",
        "Customer_Class__c",
        "Tier__c",
        "Account_Status__c",
        "Phone",
    ]
    person_fields = [
        "External_ID__c",
        "RecordTypeId",
        "FirstName",
        "LastName",
        "Institution_Type__c",
        "Customer_Class__c",
        "Tier__c",
        "Account_Status__c",
        "BillingCountry",
        "PersonMobilePhone",
        "PersonEmail",
    ]
    aff_fields = [
        "External_ID__c",
        "Primary_Account__r.External_ID__c",
        "Related_Account__r.External_ID__c",
        "Affiliation_Type__c",
        "Role__c",
        "Is_Active__c",
    ]

    write_csv(OUT / "accounts_org.csv", org_fields, org_rows)
    write_csv(OUT / "accounts_person.csv", person_fields, person_rows)
    write_csv(OUT / "affiliations.csv", aff_fields, aff_rows)
    write_csv(
        OUT / "workplace_skipped.csv",
        ["code", "doctor", "customerClass"],
        workplace_skipped,
    )

    print("Reading target…")
    target = read_sheet1(TARGET_XLSX)
    print(f"  {len(target)} rows")

    # Persist raw target for load-targets stage (needs live Account/User maps).
    target_slim = []
    for rec in target:
        target_slim.append(
            {
                "Target": norm(rec.get("Target")),
                "Code": norm(rec.get("Code")),
                "Doctor Code": norm(rec.get("Doctor Code")),
                "Doctor Clinic ID": norm(rec.get("Doctor Clinic ID")),
                "Customer Type": norm(rec.get("Customer Type")),
                "Customer Name": norm(rec.get("Customer Name")),
                "New Frequency": norm(rec.get("New Frequency")),
                "Default Frequency": norm(rec.get("Default Frequency")),
                "New Class": norm(rec.get("New Class")),
                "Employee Code": norm(rec.get("Employee Code")),
                "Employee Name": norm(rec.get("Employee Name")),
                "Team Name": norm(rec.get("Team Name")),
                "Team Code": norm(rec.get("Team Code")),
            }
        )
    write_csv(
        OUT / "target_slim.csv",
        list(target_slim[0].keys()) if target_slim else ["Code"],
        target_slim,
    )

    meta = {
        "orgAccounts": len(org_rows),
        "personAccounts": len(person_rows),
        "affiliations": len(aff_rows),
        "workplaceSkipped": len(workplace_skipped),
        "targetRows": len(target_slim),
        "expectedAccounts": len(org_rows) + len(person_rows),
    }
    (OUT / "prepare_meta.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
    print(json.dumps(meta, indent=2))
    print(f"Wrote CSVs under {OUT}")


def query_all(soql: str) -> list[dict]:
    data = sf_json(["data", "query", "--query", soql, "--result-format", "json"])
    result = data.get("result") or {}
    records = list(result.get("records") or [])
    # Follow nextRecordsUrl via --query-more if needed — for large sets use bulk query.
    return records


def bulk_query(soql: str, out_csv: Path) -> list[dict]:
    """Export all matching rows via Bulk API 2.0 (`sf data export bulk`)."""
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    if out_csv.exists():
        out_csv.unlink()
    proc = subprocess.run(
        [
            "sf",
            "data",
            "export",
            "bulk",
            "--query",
            soql,
            "--output-file",
            str(out_csv),
            "--wait",
            "60",
            "--line-ending",
            "LF",
        ],
        cwd=str(ROOT),
        capture_output=True,
        text=True,
    )
    if proc.returncode != 0 or not out_csv.exists():
        print("export bulk failed:", (proc.stderr or proc.stdout)[:800])
        # Paginate with regular query as fallback (Id > lastId).
        rows: list[dict] = []
        after = None
        while True:
            page_soql = soql
            if " ORDER BY " not in soql.upper():
                page_soql = soql + " ORDER BY Id"
            if after:
                if " WHERE " in page_soql.upper():
                    page_soql = page_soql.replace(" WHERE ", f" WHERE Id > '{after}' AND ", 1)
                    # if replace failed due to case:
                    if f"Id > '{after}'" not in page_soql:
                        idx = page_soql.upper().index("WHERE")
                        page_soql = (
                            page_soql[: idx + 5]
                            + f" Id > '{after}' AND"
                            + page_soql[idx + 5 :]
                        )
                else:
                    # INSERT WHERE before ORDER BY
                    oi = page_soql.upper().index("ORDER BY")
                    page_soql = page_soql[:oi] + f" WHERE Id > '{after}' " + page_soql[oi:]
            page_soql = page_soql + " LIMIT 2000"
            data = sf_json(["data", "query", "--query", page_soql])
            batch = (data.get("result") or {}).get("records") or []
            if not batch:
                break
            for rec in batch:
                rows.append({k: v for k, v in rec.items() if k != "attributes"})
            after = batch[-1]["Id"]
            print(f"  paged {len(rows)}…")
            if len(batch) < 2000:
                break
        if rows:
            keys = list(rows[0].keys())
            write_csv(out_csv, keys, rows)
            return rows
        raise RuntimeError(proc.stderr or proc.stdout or "bulk export failed")
    rows = []
    with out_csv.open(encoding="utf-8") as f:
        for row in csv.DictReader(f):
            rows.append(row)
    return rows


def load_accounts() -> None:
    org = OUT / "accounts_org.csv"
    person = OUT / "accounts_person.csv"
    if not org.exists() or not person.exists():
        raise SystemExit("Run prepare first")
    upsert_csv("Account", "External_ID__c", org)
    upsert_csv("Account", "External_ID__c", person)


def load_affiliations() -> None:
    path = OUT / "affiliations.csv"
    if not path.exists():
        raise SystemExit("Run prepare first")
    upsert_csv("Account_Affiliation__c", "External_ID__c", path)


def resolve_primary_territories(emp_codes: set[str]) -> tuple[dict[str, str], dict[str, dict], list[str]]:
    """Return emp→territoryId, emp→user info, and missing employees."""
    if not emp_codes:
        return {}, {}, []
    # Chunk IN clauses
    codes = sorted(emp_codes)
    users: dict[str, dict] = {}
    for i in range(0, len(codes), 200):
        chunk = codes[i : i + 200]
        in_list = ",".join("'" + c.replace("'", "\\'") + "'" for c in chunk)
        soql = (
            "SELECT Id, Name, EmployeeNumber, IsActive FROM User "
            f"WHERE EmployeeNumber IN ({in_list})"
        )
        for rec in query_all(soql):
            emp = norm(rec.get("EmployeeNumber"))
            if emp:
                users[emp] = rec
    missing = [c for c in codes if c not in users]
    user_ids = [u["Id"] for u in users.values()]
    # Prefer role containing 'rep'
    territory_by_user: dict[str, tuple[str, str]] = {}
    for i in range(0, len(user_ids), 200):
        chunk = user_ids[i : i + 200]
        in_list = ",".join("'" + c + "'" for c in chunk)
        soql = (
            "SELECT UserId, Territory2Id, RoleInTerritory2 FROM UserTerritory2Association "
            f"WHERE UserId IN ({in_list})"
        )
        for rec in query_all(soql):
            uid = rec["UserId"]
            role = norm(rec.get("RoleInTerritory2"))
            cand_rep = "rep" in role.lower()
            if uid not in territory_by_user:
                territory_by_user[uid] = (rec["Territory2Id"], role)
            else:
                _, existing_role = territory_by_user[uid]
                exist_rep = "rep" in existing_role.lower()
                if cand_rep and not exist_rep:
                    territory_by_user[uid] = (rec["Territory2Id"], role)

    emp_to_terr: dict[str, str] = {}
    for emp, u in users.items():
        pair = territory_by_user.get(u["Id"])
        if pair:
            emp_to_terr[emp] = pair[0]
    return emp_to_terr, users, missing


def load_account_id_map() -> dict[str, str]:
    out_csv = OUT / "account_ext_ids.csv"
    print("Bulk querying Account External_ID__c…")
    rows = bulk_query("SELECT Id, External_ID__c FROM Account WHERE External_ID__c != null", out_csv)
    mapping = {}
    for r in rows:
        ext = norm(r.get("External_ID__c"))
        if ext:
            mapping[ext] = r["Id"]
    print(f"  {len(mapping)} accounts with External_ID__c")
    return mapping


def resolve_target_ext(row: dict, account_by_ext: dict[str, str], workplace_docs: set[str]) -> str | None:
    """Deterministic match: Doctor → DOC- then CLINIC-; org → Code only."""
    ctype = norm(row.get("Customer Type")).lower()
    code = norm(row.get("Code"))
    doctor_code = norm(row.get("Doctor Code"))
    clinic_id = norm(row.get("Doctor Clinic ID"))

    if ctype == "doctor":
        if doctor_code and doctor_code not in workplace_docs:
            ext = doctor_ext(doctor_code)
            if ext in account_by_ext:
                return ext
        if clinic_id:
            ext = clinic_ext(clinic_id)
            if ext in account_by_ext:
                return ext
        return None

    if code and code in account_by_ext:
        return code
    return None


def load_targets() -> None:
    slim_path = OUT / "target_slim.csv"
    if not slim_path.exists():
        raise SystemExit("Run prepare first")

    workplace_docs = set()
    wp = OUT / "workplace_skipped.csv"
    if wp.exists():
        with wp.open(encoding="utf-8") as f:
            for row in csv.DictReader(f):
                workplace_docs.add(norm(row.get("code")))

    account_by_ext = load_account_id_map()

    target_rows = []
    with slim_path.open(encoding="utf-8") as f:
        target_rows = list(csv.DictReader(f))

    emp_codes = set()
    for row in target_rows:
        emp = norm(row.get("Employee Code"))
        if emp and not emp.lower().startswith("vac"):
            emp_codes.add(emp)

    print(f"Resolving territories for {len(emp_codes)} employees…")
    emp_to_terr, users, missing_emps = resolve_primary_territories(emp_codes)
    write_csv(
        OUT / "missing_employees.csv",
        ["Employee Code"],
        [{"Employee Code": e} for e in missing_emps],
    )
    no_terr = [e for e in emp_codes if e in users and e not in emp_to_terr]
    write_csv(
        OUT / "employees_without_territory.csv",
        ["Employee Code", "Name"],
        [{"Employee Code": e, "Name": users[e].get("Name")} for e in no_terr],
    )

    # Territory line name for team mismatch report
    terr_ids = sorted(set(emp_to_terr.values()))
    terr_name: dict[str, str] = {}
    for i in range(0, len(terr_ids), 200):
        chunk = terr_ids[i : i + 200]
        in_list = ",".join("'" + c + "'" for c in chunk)
        soql = (
            "SELECT Id, Name, ParentTerritory2.Name, ParentTerritory2.ParentTerritory2.Name "
            f"FROM Territory2 WHERE Id IN ({in_list})"
        )
        for rec in query_all(soql):
            # Prefer line (grandparent of position often): Position → District → Line
            line = None
            parent = rec.get("ParentTerritory2") or {}
            if isinstance(parent, dict):
                gp = parent.get("ParentTerritory2") or {}
                if isinstance(gp, dict):
                    line = gp.get("Name")
                if not line:
                    line = parent.get("Name")
            terr_name[rec["Id"]] = line or rec.get("Name") or ""

    atf_by_key: dict[str, dict] = {}
    account_class_updates: dict[str, dict] = {}
    exceptions: list[dict] = []
    multi_clinic_pairs: list[dict] = []
    team_mismatches: list[dict] = []
    seen_doc_emp_clinics: dict[tuple[str, str], list[tuple[str, str, str]]] = defaultdict(list)

    for row in target_rows:
        emp = norm(row.get("Employee Code"))
        if not emp or emp.lower().startswith("vac"):
            exceptions.append({**row, "Reason": "vacant"})
            continue
        if emp not in users:
            exceptions.append({**row, "Reason": "no_user"})
            continue
        terr = emp_to_terr.get(emp)
        if not terr:
            exceptions.append({**row, "Reason": "no_territory"})
            continue

        team = norm(row.get("Team Name"))
        line = terr_name.get(terr, "")
        if team and line and team.upper() not in line.upper() and line.upper() not in team.upper():
            team_mismatches.append(
                {
                    "Employee Code": emp,
                    "Employee Name": users[emp].get("Name"),
                    "Team Name": team,
                    "Territory Line": line,
                    "Territory2Id": terr,
                }
            )

        ext = resolve_target_ext(row, account_by_ext, workplace_docs)
        if not ext:
            exceptions.append({**row, "Reason": "no_account"})
            continue

        account_id = account_by_ext[ext]
        key = (account_id + "_" + terr)[:40]
        freq_raw = norm(row.get("New Frequency")) or norm(row.get("Default Frequency"))
        try:
            visits = float(freq_raw) if freq_raw != "" else None
        except ValueError:
            visits = None
        cls = norm(row.get("New Class"))
        tier = map_tier(cls)

        ctype = norm(row.get("Customer Type")).lower()
        if ctype == "doctor":
            dcode = norm(row.get("Doctor Code"))
            cid = norm(row.get("Doctor Clinic ID"))
            seen_doc_emp_clinics[(dcode, emp)].append((cid, cls, freq_raw))

        existing = atf_by_key.get(key)
        if existing:
            # Keep higher visit frequency when same account+territory repeats.
            prev = existing.get("Call_Plan_Visits_Per_Month__c")
            if visits is not None and (prev == "" or prev is None or float(prev) < visits):
                existing["Call_Plan_Visits_Per_Month__c"] = visits
            if cls and not existing.get("HCP_Classification__c"):
                existing["HCP_Classification__c"] = cls[:10]
                existing["Classification__c"] = tier
                existing["Potential__c"] = tier
        else:
            atf_by_key[key] = {
                "Unique_Key__c": key,
                "Account__c": account_id,
                "Territory2_Id__c": terr,
                "Is_Active__c": "true",
                "Classification__c": tier or "",
                "HCP_Classification__c": (cls[:10] if cls else ""),
                "Potential__c": tier or "",
                "Call_Plan_Visits_Per_Month__c": "" if visits is None else visits,
            }

        if cls:
            account_class_updates[account_id] = {
                "Id": account_id,
                "Customer_Class__c": cls[:255],
                "Tier__c": tier or "",
            }

    for (dcode, emp), clinics in seen_doc_emp_clinics.items():
        unique_clinics = {c[0] for c in clinics}
        if len(unique_clinics) > 1:
            multi_clinic_pairs.append(
                {
                    "Doctor Code": dcode,
                    "Employee Code": emp,
                    "Clinics": ";".join(sorted(unique_clinics)),
                    "Classes": ";".join(sorted({c[1] for c in clinics})),
                    "Frequencies": ";".join(sorted({c[2] for c in clinics})),
                }
            )

    # Deduplicate team mismatches
    tm_by = {(r["Employee Code"], r["Team Name"], r["Territory Line"]): r for r in team_mismatches}
    team_mismatches = list(tm_by.values())

    atf_rows = list(atf_by_key.values())
    atf_fields = [
        "Unique_Key__c",
        "Account__c",
        "Territory2_Id__c",
        "Is_Active__c",
        "Classification__c",
        "HCP_Classification__c",
        "Potential__c",
        "Call_Plan_Visits_Per_Month__c",
    ]
    write_csv(OUT / "atf.csv", atf_fields, atf_rows)
    write_csv(
        OUT / "account_class_updates.csv",
        ["Id", "Customer_Class__c", "Tier__c"],
        list(account_class_updates.values()),
    )
    write_csv(
        OUT / "target_exceptions.csv",
        list(target_rows[0].keys()) + ["Reason"] if target_rows else ["Reason"],
        exceptions,
    )
    write_csv(
        OUT / "multi_clinic_same_rep.csv",
        ["Doctor Code", "Employee Code", "Clinics", "Classes", "Frequencies"],
        multi_clinic_pairs,
    )
    write_csv(
        OUT / "team_vs_territory_mismatches.csv",
        ["Employee Code", "Employee Name", "Team Name", "Territory Line", "Territory2Id"],
        team_mismatches,
    )

    reason_counts = Counter(e.get("Reason") for e in exceptions)
    summary = {
        "atfRows": len(atf_rows),
        "accountClassUpdates": len(account_class_updates),
        "exceptions": len(exceptions),
        "exceptionReasons": dict(reason_counts),
        "multiClinicSameRep": len(multi_clinic_pairs),
        "teamMismatches": len(team_mismatches),
        "missingEmployees": len(missing_emps),
        "employeesWithoutTerritory": len(no_terr),
    }
    (OUT / "target_load_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))

    upsert_csv("Account_Territory_Fields__c", "Unique_Key__c", OUT / "atf.csv")

    # Account class/tier update via bulk upsert on Id (update).
    class_csv = OUT / "account_class_updates.csv"
    if account_class_updates:
        print("Updating Account Customer_Class__c / Tier__c…")
        data = sf_json(
            [
                "data",
                "update",
                "bulk",
                "--sobject",
                "Account",
                "--file",
                str(class_csv),
                "--line-ending",
                "LF",
                "--wait",
                "180",
            ]
        )
        result = data.get("result") or {}
        print(
            "  update job=",
            result.get("jobId") or result.get("id"),
            "failed=",
            result.get("numberRecordsFailed"),
        )


def report() -> None:
    counts = {}
    for label, soql in [
        ("Account", "SELECT COUNT() FROM Account"),
        ("Affiliation", "SELECT COUNT() FROM Account_Affiliation__c"),
        ("ATF", "SELECT COUNT() FROM Account_Territory_Fields__c"),
        ("Brick", "SELECT COUNT() FROM Brick__c"),
        ("Territory2", "SELECT COUNT() FROM Territory2"),
    ]:
        data = sf_json(["data", "query", "--query", soql])
        counts[label] = (data.get("result") or {}).get("totalSize")
    meta = {}
    if (OUT / "prepare_meta.json").exists():
        meta = json.loads((OUT / "prepare_meta.json").read_text())
    target_sum = {}
    if (OUT / "target_load_summary.json").exists():
        target_sum = json.loads((OUT / "target_load_summary.json").read_text())
    print(json.dumps({"orgCounts": counts, "prepare": meta, "target": target_sum}, indent=2))


def wait_wipe_zero(poll_s: int = 60) -> None:
    while True:
        data = sf_json(["data", "query", "--query", "SELECT COUNT() FROM Account"])
        n = (data.get("result") or {}).get("totalSize")
        jobs = sf_json(
            [
                "data",
                "query",
                "--query",
                "SELECT Id, Status FROM AsyncApexJob WHERE ApexClass.Name='AccountWipeBatch' "
                "AND Status IN ('Queued','Preparing','Processing','Holding')",
            ]
        )
        active = (jobs.get("result") or {}).get("totalSize") or 0
        print(f"Accounts={n} activeWipeJobs={active}")
        if n == 0 and active == 0:
            return
        time.sleep(poll_s)


def run_all() -> None:
    prepare()
    print("Waiting for wipe to reach 0 accounts…")
    wait_wipe_zero()
    load_accounts()
    load_affiliations()
    load_targets()
    report()


def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    cmd = sys.argv[1]
    {
        "prepare": prepare,
        "load-accounts": load_accounts,
        "load-affiliations": load_affiliations,
        "load-targets": load_targets,
        "report": report,
        "wait-wipe": wait_wipe_zero,
        "all": run_all,
    }[cmd]()


if __name__ == "__main__":
    main()
