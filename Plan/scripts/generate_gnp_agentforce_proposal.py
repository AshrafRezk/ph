#!/usr/bin/env python3
"""Generate the GNP Agentforce Life Sciences proposal (editable Word document)."""

from datetime import date
from pathlib import Path
from typing import List, Optional, Sequence, Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "GNP_Agentforce_Life_Sciences_Proposal.docx"
SF_LOGO = ROOT.parent / "offline" / "public" / "salesforce-logo-512.png"

# Salesforce brand colours (official)
NAVY = RGBColor(0x03, 0x2D, 0x60)
BLUE = RGBColor(0x01, 0x76, 0xD3)
CLOUD = RGBColor(0x0D, 0x9D, 0xDA)
INK = RGBColor(0x18, 0x18, 0x18)
GRAY = RGBColor(0x44, 0x44, 0x44)
MUTED = RGBColor(0x70, 0x6E, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

NAVY_HEX = "032D60"
BLUE_HEX = "0176D3"
LIGHT_HEX = "EAF5FE"
SOFT_HEX = "F4F6F9"
ROW_ALT_HEX = "F3F8FD"


def set_run_font(run, *, name="Calibri", size=11, bold=False, italic=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:cs"), name)


def set_paragraph_spacing(paragraph, before=0, after=8, line=15, line_rule=WD_LINE_SPACING.MULTIPLE):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if line_rule == WD_LINE_SPACING.EXACTLY:
        fmt.line_spacing = Pt(line)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    else:
        fmt.line_spacing = line / 12.0
        fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def shade_cell(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_borders(cell, color="D8DDE6", sz="4") -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def prevent_row_split(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def set_table_width(table, width_inches: float) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(width_inches * 1440)))
    tblW.set(qn("w:type"), "dxa")


def disable_autofit(table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def add_bottom_border(paragraph, color="0176D3", sz="12") -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    set_run_font(run, size=8, color=MUTED)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color, before, after, bold in (
        ("Heading 1", 16, NAVY, 16, 8, True),
        ("Heading 2", 13, BLUE, 14, 6, True),
        ("Heading 3", 11.5, NAVY, 10, 4, True),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = color
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_body(doc: Document, text: str, *, size=11, color=INK, italic=False, bold=False, after=8, before=0):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, italic=italic, bold=bold)
    return p


def add_rich(doc: Document, parts: Sequence[Tuple], *, after=8, before=0):
    """parts: (text, bold, italic, color, size) tuples."""
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after)
    for item in parts:
        text = item[0]
        bold = item[1] if len(item) > 1 else False
        italic = item[2] if len(item) > 2 else False
        color = item[3] if len(item) > 3 else INK
        size = item[4] if len(item) > 4 else 11
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_bullet(doc: Document, text: str, *, bold_lead: Optional[str] = None):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, before=0, after=4)
    if bold_lead:
        lead = p.add_run(bold_lead)
        set_run_font(lead, size=11, bold=True, color=NAVY)
        rest = p.add_run(text)
        set_run_font(rest, size=11, color=INK)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=INK)
    return p


def add_callout(doc: Document, title: str, body: str, *, fill=LIGHT_HEX, accent=BLUE_HEX):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    disable_autofit(table)
    set_table_width(table, 6.7)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_borders(cell, color=accent, sz="8")
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    cell.width = Inches(6.7)
    cell.text = ""
    p1 = cell.paragraphs[0]
    set_paragraph_spacing(p1, after=4)
    r1 = p1.add_run(title)
    set_run_font(r1, size=10, bold=True, color=NAVY)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, after=0)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10, color=GRAY)
    doc.add_paragraph()
    return table


def add_table(
    doc: Document,
    headers: List[str],
    rows: List[Tuple],
    col_widths: Optional[List[float]] = None,
    header_fill: str = NAVY_HEX,
    header_size: float = 9.5,
    body_size: float = 9.5,
    bold_rows: Optional[List[int]] = None,
) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    disable_autofit(table)
    bold_rows = bold_rows or []
    total = sum(col_widths) if col_widths else 6.7
    set_table_width(table, total)

    hdr = table.rows[0]
    prevent_row_split(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        shade_cell(cell, header_fill)
        set_cell_borders(cell, color=header_fill, sz="4")
        set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
        cell.text = ""
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, after=0, line=12, line_rule=WD_LINE_SPACING.EXACTLY)
        run = p.add_run(header)
        set_run_font(run, size=header_size, bold=True, color=WHITE)
        if col_widths:
            cell.width = Inches(col_widths[i])

    for r_idx, row_data in enumerate(rows, start=1):
        row = table.rows[r_idx]
        prevent_row_split(row)
        fill = ROW_ALT_HEX if r_idx % 2 == 0 else "FFFFFF"
        is_bold = r_idx in bold_rows
        for c_idx, value in enumerate(row_data):
            cell = row.cells[c_idx]
            shade_cell(cell, fill)
            set_cell_borders(cell, color="D8DDE6", sz="4")
            set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
            cell.text = ""
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, after=0, line=13, line_rule=WD_LINE_SPACING.EXACTLY)
            run = p.add_run(str(value))
            set_run_font(run, size=body_size, bold=is_bold, color=INK)
            if col_widths:
                cell.width = Inches(col_widths[c_idx])

    doc.add_paragraph()


def configure_section(section, *, different_first=False) -> None:
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    if different_first:
        section.different_first_page_header_footer = True


def add_header_footer(section, *, first_page=False) -> None:
    header = section.first_page_header if first_page else section.header
    footer = section.first_page_footer if first_page else section.footer
    header.is_linked_to_previous = False
    footer.is_linked_to_previous = False

    if first_page:
        header.paragraphs[0].text = ""
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.text = ""
        run = fp.add_run("Confidential  ·  For Global Napi Pharmaceuticals  ·  Not for onward distribution")
        set_run_font(run, size=8, color=MUTED, italic=True)
        return

    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(hp, after=2)
    add_bottom_border(hp, color=BLUE_HEX, sz="18")
    r1 = hp.add_run("Salesforce  ")
    set_run_font(r1, size=9, bold=True, color=NAVY)
    r2 = hp.add_run("Agentforce Life Sciences")
    set_run_font(r2, size=9, color=BLUE)
    tab = hp.add_run("\t")
    set_run_font(tab, size=9)
    r3 = hp.add_run("Cloudastick  ·  Salesforce Partner")
    set_run_font(r3, size=8.5, color=MUTED)
    # right-align via tabs
    pPr = hp._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab_el = OxmlElement("w:tab")
    tab_el.set(qn("w:val"), "right")
    tab_el.set(qn("w:pos"), "8640")
    tabs.append(tab_el)
    pPr.append(tabs)

    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(fp, before=2)
    pPrf = fp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "6")
    top.set(qn("w:space"), "6")
    top.set(qn("w:color"), "D8DDE6")
    pBdr.append(top)
    pPrf.append(pBdr)
    r4 = fp.add_run("Prepared for GNP  ·  Cloudastick Systems  ·  Salesforce Partner\t")
    set_run_font(r4, size=8, color=MUTED)
    add_page_number(fp)
    tabsf = OxmlElement("w:tabs")
    tabf = OxmlElement("w:tab")
    tabf.set(qn("w:val"), "right")
    tabf.set(qn("w:pos"), "8640")
    tabsf.append(tabf)
    pPrf.append(tabsf)


def add_cover(doc: Document) -> None:
    banner = doc.add_table(rows=1, cols=1)
    disable_autofit(banner)
    set_table_width(banner, 6.7)
    cell = banner.cell(0, 0)
    shade_cell(cell, NAVY_HEX)
    set_cell_borders(cell, color=NAVY_HEX, sz="0")
    set_cell_margins(cell, top=160, bottom=160, left=140, right=140)
    cell.text = ""

    p_logo = cell.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p_logo, after=6)
    if SF_LOGO.exists():
        run = p_logo.add_run()
        run.add_picture(str(SF_LOGO), width=Inches(0.55))

    p_kicker = cell.add_paragraph()
    set_paragraph_spacing(p_kicker, before=8, after=2)
    rk = p_kicker.add_run("SALESFORCE PARTNER PROPOSAL")
    set_run_font(rk, size=9, bold=True, color=CLOUD)

    p_title = cell.add_paragraph()
    set_paragraph_spacing(p_title, before=4, after=4)
    rt = p_title.add_run("Agentforce Life Sciences")
    set_run_font(rt, size=26, bold=True, color=WHITE)

    p_sub = cell.add_paragraph()
    set_paragraph_spacing(p_sub, after=2)
    rs = p_sub.add_run("for Global Napi Pharmaceuticals")
    set_run_font(rs, size=16, color=CLOUD)

    p_tag = cell.add_paragraph()
    set_paragraph_spacing(p_tag, before=8, after=0)
    rg = p_tag.add_run(
        "Adding trusted physician master data and AI insights\n"
        "to GNP’s existing commercial platform — without replacing it."
    )
    set_run_font(rg, size=11, color=WHITE)

    doc.add_paragraph()

    meta = doc.add_table(rows=4, cols=2)
    disable_autofit(meta)
    set_table_width(meta, 6.7)
    meta_rows = [
        ("Prepared for", "Global Napi Pharmaceuticals (GNP)"),
        ("Prepared by", "Cloudastick Systems  ·  Salesforce Partner"),
        ("Solution", "Salesforce Agentforce Life Sciences + Cloudastick Intelligence Layer"),
        ("Date", f"{date.today():%B %Y}  ·  Discussion document  ·  Editable"),
    ]
    for i, (label, value) in enumerate(meta_rows):
        prevent_row_split(meta.rows[i])
        lc, vc = meta.rows[i].cells
        shade_cell(lc, SOFT_HEX)
        shade_cell(vc, "FFFFFF")
        set_cell_borders(lc, color="E3E8EF")
        set_cell_borders(vc, color="E3E8EF")
        set_cell_margins(lc, top=50, bottom=50, left=80, right=80)
        set_cell_margins(vc, top=50, bottom=50, left=80, right=80)
        lc.text = ""
        vc.text = ""
        lp = lc.paragraphs[0]
        vp = vc.paragraphs[0]
        set_paragraph_spacing(lp, after=0)
        set_paragraph_spacing(vp, after=0)
        lr = lp.add_run(label)
        vr = vp.add_run(value)
        set_run_font(lr, size=10, bold=True, color=NAVY)
        set_run_font(vr, size=10, color=INK)
        lc.width = Inches(1.8)
        vc.width = Inches(4.9)

    doc.add_paragraph()
    add_callout(
        doc,
        "How to read this document",
        "This is a conservative proposal for discussion. It describes how Salesforce Agentforce "
        "Life Sciences and Cloudastick’s pharmaceutical intelligence layer can sit beside GNP’s "
        "in-house system. Nothing here replaces GNP’s tracking or physician reporting. Scope, "
        "interfaces, and commercials will be confirmed after a short discovery. The Word file is "
        "fully editable — headings, tables, and wording can be adjusted before sharing internally.",
        fill=LIGHT_HEX,
        accent=BLUE_HEX,
    )


def add_contents(doc: Document) -> None:
    doc.add_heading("Contents", level=1)
    items = [
        "1.  What we heard from GNP",
        "2.  Guiding principle — add, do not rebuild",
        "3.  Proposed architecture",
        "4.  Search Before Create — physician, pharmacy, and HCO master",
        "5.  Agentforce intelligence layer",
        "     5.1  Physician Intelligence",
        "     5.2  Visit Insights",
        "     5.3  Absolute Physician Movement",
        "     5.4  Distributor sales data and product-code unification",
        "     5.5  Suggestions and next-best action",
        "6.  Merging GNP’s AI with Cloudastick Agentforce",
        "7.  What stays with GNP, what we add",
        "8.  Indicative delivery approach",
        "9.  Trust, compliance, and data use",
        "10. Suggested next step",
        "Appendix  —  Agent catalogue",
    ]
    for item in items:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=3)
        run = p.add_run(item)
        set_run_font(run, size=11, color=NAVY)


def add_heard(doc: Document) -> None:
    doc.add_heading("1.  What we heard from GNP", level=1)
    add_body(
        doc,
        "GNP already runs an in-house commercial platform. Field tracking and physician reporting "
        "are in place. The conversation is not about discarding that investment. It is about "
        "giving those records a trusted customer master they do not have today, and giving "
        "managers and reps Agentforce insights on top of work they already capture.",
    )
    add_table(
        doc,
        ["GNP position", "What this proposal answers"],
        [
            (
                "Keep the in-house system",
                "Salesforce and Cloudastick sit beside GNP as an intelligence and master-data layer, connected by APIs. GNP remains system of record for visits and reporting.",
            ),
            (
                "Need Salesforce AI — suggestions, not a new CRM",
                "Agentforce Life Sciences produces grounded suggestions and briefs from GNP data plus approved public professional context.",
            ),
            (
                "Tracking and physician reporting exist; insights do not",
                "Visit Insights, Absolute Physician Movement, and next-best action read GNP’s existing activity instead of asking reps to recapture it.",
            ),
            (
                "Merge our AI with theirs",
                "GNP’s AI module and Cloudastick’s Agentforce skills share one agent contract. They call each other; they do not compete as two chatbots.",
            ),
            (
                "No physician database of their own",
                "Cloudastick Search Before Create looks up a live-updated HCP, pharmacy, and HCO master before any new record is added to GNP.",
            ),
        ],
        col_widths=[2.4, 4.3],
    )


def add_principle(doc: Document) -> None:
    doc.add_heading("2.  Guiding principle — add, do not rebuild", level=1)
    add_body(
        doc,
        "Pharmaceutical field platforms fail when a vendor asks the company to start again. "
        "GNP’s users already know their screens. Their visit history is an asset. The right "
        "Salesforce motion is to extend that asset with Agentforce, not to migrate the field "
        "force onto a green-field CRM on day one.",
    )
    add_callout(
        doc,
        "Working rule",
        "If GNP already captures it, we read it. If GNP already reports it, we interpret it. "
        "If GNP is missing a physician, pharmacy, or HCO, we search Cloudastick’s live master "
        "before create. New Salesforce screens are introduced only where they add intelligence "
        "GNP cannot produce today.",
    )
    add_body(doc, "In practice that means four commitments:", before=2)
    add_bullet(
        doc,
        " GNP’s in-house application remains the daily workspace unless GNP later chooses a deeper Salesforce field shell.",
        bold_lead="System of record stays with GNP.",
    )
    add_bullet(
        doc,
        " Visits, physician reports, and existing AI outputs are ingested or queried — they are not re-keyed.",
        bold_lead="No double entry.",
    )
    add_bullet(
        doc,
        " Search Before Create, Agentforce insights, and master enrichment can be exposed inside GNP’s UI via API, or in a Salesforce console for SFE and data stewards.",
        bold_lead="Embed or sidecar.",
    )
    add_bullet(
        doc,
        " Discovery confirms interfaces, identity, and which Agentforce skills go live first. This document is the map, not the signed scope.",
        bold_lead="Prove value in phases.",
    )


def add_architecture(doc: Document) -> None:
    doc.add_heading("3.  Proposed architecture", level=1)
    add_body(
        doc,
        "The design is a four-layer stack. GNP does not move data into Salesforce in order to "
        "“have a CRM.” Salesforce Agentforce Life Sciences is introduced because it is the "
        "trusted place to run life-sciences agents, with Data 360 for harmonisation and the "
        "Einstein Trust Layer for audit.",
    )
    add_table(
        doc,
        ["Layer", "Owner", "Role for GNP"],
        [
            (
                "GNP in-house platform",
                "GNP",
                "System of record for visits, tracking, physician reporting, and GNP’s existing AI module.",
            ),
            (
                "Integration fabric",
                "Joint",
                "Secure APIs, events, and Data 360 connectors. Account, visit, product, and territory identity are mapped once.",
            ),
            (
                "Agentforce Life Sciences",
                "Salesforce + Cloudastick",
                "Agent orchestration, grounded suggestions, visit and account briefs, apply-actions, Einstein Trust Layer.",
            ),
            (
                "Cloudastick intelligence",
                "Cloudastick",
                "Live-updated HCP / pharmacy / HCO master, Search Before Create, Map Intelligence, pharma-specific Agentforce skills.",
            ),
        ],
        col_widths=[1.85, 1.45, 3.4],
    )
    add_body(
        doc,
        "GNP users can keep working in the in-house application. Agentforce results return as "
        "cards, briefs, and recommended actions on those same screens. Data stewards and SFE "
        "may additionally use a Salesforce console for master matching, product-code mapping, "
        "and insight review.",
    )
    add_body(
        doc,
        "This is consistent with how Agentforce Life Sciences is used across the industry: "
        "agents sit on a unified data foundation and take action in the systems commercial "
        "teams already run — they are not a second field-force product that has to win the "
        "desktop first.",
        italic=True,
        size=10.5,
        color=GRAY,
    )


def add_search_before_create(doc: Document) -> None:
    doc.add_heading("4.  Search Before Create — physician, pharmacy, and HCO master", level=1)
    add_body(
        doc,
        "GNP does not have a physician database of its own. Every new doctor, pharmacy, or "
        "healthcare organisation added from the field therefore risks becoming a duplicate, "
        "an incomplete stub, or a record that never joins later sell-out and visit history. "
        "Search Before Create is the Cloudastick service that closes that gap without asking "
        "GNP to build a master from scratch.",
    )
    doc.add_heading("How it works in GNP’s create flow", level=2)
    add_body(
        doc,
        "Before a new physician, pharmacy, or HCO is written into GNP’s system, the create "
        "step calls Cloudastick’s live-updated master. The master is searched by name, "
        "specialty, brick, city, phone, and related identifiers. Cognitive matching handles "
        "Arabic and Latin spelling variants so “Ahmed” and “أحمد” can resolve to the same person.",
    )
    add_table(
        doc,
        ["Step", "What happens", "Outcome in GNP"],
        [
            (
                "1. Search master",
                "Rep or data steward searches Cloudastick’s live HCP / pharmacy / HCO list, with brick and specialty filters.",
                "Existing master record is proposed. No new row is created yet.",
            ),
            (
                "2. Adopt or enrich",
                "If a confident match exists, GNP stores the Cloudastick identity (and agreed profile fields) on the local record.",
                "One physician, one identity — usable for visits and reporting immediately or after GNP’s approval rule.",
            ),
            (
                "3. Map Intelligence",
                "If the master has no match, Cloudastick Map Intelligence searches public place and professional location signals across Egypt (zone, brick, area).",
                "A candidate place can seed a quality-scored create form instead of a blank page.",
            ),
            (
                "4. Controlled create",
                "Only then may a genuinely new record be added, with a data-quality score and optional approval.",
                "GNP’s database grows without silent duplicates. The master continues to update live.",
            ),
        ],
        col_widths=[1.5, 2.7, 2.5],
    )
    add_callout(
        doc,
        "Why this is the first capability",
        "Agentforce cannot give reliable physician intelligence on a noisy customer list. "
        "Search Before Create is the conservative foundation: GNP keeps its system, and "
        "gains a live master it does not have to maintain alone. Every later insight — "
        "visit briefs, movement, sell-out — hangs off that identity.",
    )


def add_agentforce_layer(doc: Document) -> None:
    doc.add_heading("5.  Agentforce intelligence layer", level=1)
    add_body(
        doc,
        "GNP already tracks activity and reports on physicians. The missing piece is an "
        "agent that reads those records, adds master and market context, and returns a "
        "suggestion a person can accept. The skills below are Cloudastick pharmaceutical "
        "intelligence, presented and governed as Salesforce Agentforce Life Sciences — "
        "not a separate chatbot brand, and not a replacement reporting suite.",
    )

    doc.add_heading("5.1  Physician Intelligence", level=2)
    add_body(
        doc,
        "Agentforce Physician Intelligence builds a working picture of an HCP or organisation "
        "from three sources GNP can trust: GNP’s own visits and ratings, Cloudastick’s master "
        "profile, and publicly available professional context (hospital or clinic association, "
        "specialty signals, and similar open professional information).",
    )
    add_bullet(doc, " Grounded account briefs for the next call — what was discussed, what is overdue, who is affiliated.")
    add_bullet(doc, " Rating validity — Agentforce checks whether the rep’s classification of an HCP or HCO is consistent with visit behaviour and public professional signals.")
    add_bullet(doc, " Suggestions, not silent overwrites — the agent flags a mismatch; GNP’s user or steward decides.")
    add_callout(
        doc,
        "A note on “data from the internet”",
        "This is not unrestricted web scraping, and it is not a promise to invent a medical "
        "record from search results. Retrieval is limited to approved public professional "
        "sources, grounded against GNP and master data, and logged through the Einstein Trust "
        "Layer. Private, patient, or non-public material is out of scope.",
        fill=SOFT_HEX,
        accent=NAVY_HEX,
    )

    doc.add_heading("5.2  Visit Insights", level=2)
    add_body(
        doc,
        "Because GNP already captures visits, Agentforce Visit Insights can start from those "
        "records on day one of integration. The agent answers the questions reporting alone "
        "does not: am I on pace, who nearby is still unplanned, and what should I say in this call?",
    )
    add_table(
        doc,
        ["Insight", "Drawn from", "Returned to the user"],
        [
            (
                "Pre-call brief",
                "Prior visits, products discussed, ratings, affiliations",
                "A short Agentforce summary on the account or visit, in GNP or Salesforce.",
            ),
            (
                "Pace against plan",
                "GNP visit history and plan-cycle targets",
                "Ahead / on track / behind / critical — with the accounts that close the gap.",
            ),
            (
                "Nearby follow-up",
                "Geocoded accounts and remaining plan targets",
                "Suggested same-area calls the rep can accept as draft visits.",
            ),
            (
                "Team context",
                "Other GNP visits to the same HCP",
                "What a colleague already covered, so the next call does not repeat it.",
            ),
        ],
        col_widths=[1.6, 2.4, 2.7],
    )
    add_body(
        doc,
        "Where GNP wants action, not only narrative, the same agent can apply an accepted "
        "recommendation — for example creating draft visits — in GNP via API or in Salesforce "
        "if that is the planning surface agreed in discovery.",
    )

    doc.add_heading("5.3  Absolute Physician Movement", level=2)
    add_body(
        doc,
        "Absolute Physician Movement (APM) is the Agentforce view of whether an HCP actually "
        "moved on a brand — not whether the rep feels the relationship improved. It compares "
        "the physician’s classification and written potential over time with affiliated "
        "pharmacy sell-out, then ranks the account in the organisation.",
    )
    add_bullet(doc, " Monthly classification / penetration snapshots versus sell-out on the same product.")
    add_bullet(doc, " Gap detection — classification moving while withdrawals are flat, or sell-out running ahead of the rating.")
    add_bullet(doc, " Product ranking for the HCP: written, potential, gap, and organisation rank.")
    add_bullet(doc, " A plain-language Agentforce headline a manager can use in coaching, without opening five reports.")
    add_body(
        doc,
        "APM is particularly valuable for GNP because physician reporting already exists. "
        "The agent does not ask for a new rating culture; it interprets the ratings GNP has, "
        "against market movement, once product and pharmacy identity have been unified.",
        before=4,
    )

    doc.add_heading("5.4  Distributor sales data and product-code unification", level=2)
    add_body(
        doc,
        "Distributor files are commercially decisive and operationally messy. Each wholesaler "
        "brings its own item codes, client codes, and pharmacy identifiers. Without unification, "
        "Agentforce cannot tell GNP whether Empacoza grew in Maadi or whether three codes "
        "describe the same pack.",
    )
    add_body(doc, "Cloudastick’s market-unification skill, exposed through Agentforce, does three conservative jobs:")
    add_bullet(
        doc,
        " Distributor code variants, item codes, and client codes are normalised into a single product and outlet identity.",
        bold_lead="Unify product and pharmacy codes.",
    )
    add_bullet(
        doc,
        " Rows are validated before commit (known sources, valid months, positive quantities, matchable identifiers). Unmatched rows stay in a steward queue rather than polluting analytics.",
        bold_lead="Load with a quality gate.",
    )
    add_bullet(
        doc,
        " Once sell-out sits on the same accounts the field visits, Agentforce can recommend which affiliated physicians to see — turning a distributor file into a plan-cycle action.",
        bold_lead="Connect market to field.",
    )
    add_body(
        doc,
        "GNP does not need to change how distributors send files. The unification layer absorbs "
        "format differences. Which distributors and which product catalogue GNP uses is confirmed "
        "in discovery; the method does not depend on a single wholesaler.",
        italic=True,
        size=10.5,
        color=GRAY,
    )

    doc.add_heading("5.5  Suggestions and next-best action", level=2)
    add_body(
        doc,
        "Agentforce suggestions are the visible payoff of the layers above. They are ranked, "
        "explainable, and optional. A typical cycle:",
    )
    add_bullet(doc, " Sell-out or visit-pace signal appears in unified data.")
    add_bullet(doc, " Agentforce produces a short list of next-best customers or follow-up visits, with a reason.")
    add_bullet(doc, " The rep or manager accepts some, rejects others.")
    add_bullet(doc, " Accepted items become draft visits or plan updates in GNP (or Salesforce), with an audit trail.")
    add_body(
        doc,
        "This is the difference between “AI that chats” and Agentforce Life Sciences: the "
        "agent is allowed to propose records in GNP’s own process, and is not allowed to "
        "bypass it.",
        before=4,
    )


def add_merge(doc: Document) -> None:
    doc.add_heading("6.  Merging GNP’s AI with Cloudastick Agentforce", level=1)
    add_body(
        doc,
        "GNP asked that Cloudastick’s AI module merge with GNP’s AI module — Salesforce "
        "Agentforce Life Sciences integrated with the in-house platform. The conservative "
        "way to do that is a shared agent contract, not a rip-and-replace of either model.",
    )
    add_table(
        doc,
        ["Principle", "What it means"],
        [
            (
                "One contract, many skills",
                "Every agent — GNP’s or Cloudastick’s — accepts the same context (account, visit, product, territory, user) and returns insight plus optional apply-actions.",
            ),
            (
                "Agentforce as the fabric",
                "Salesforce Agentforce Life Sciences orchestrates skills, applies the Einstein Trust Layer, and records why a suggestion was made. GNP’s existing AI is registered as a peer skill, not retired.",
            ),
            (
                "GNP UI stays primary",
                "Insights can render inside GNP screens (API or embedded component). Salesforce is used where GNP prefers a steward / SFE console.",
            ),
            (
                "Grounding on GNP data",
                "Prompts are grounded on GNP visits, reports, and the Cloudastick master. The model is not trained on GNP confidential data.",
            ),
            (
                "Apply is permissioned",
                "Creating a visit, updating a plan, or writing a master identity back to GNP requires the same approvals GNP already expects.",
            ),
        ],
        col_widths=[2.1, 4.6],
    )
    add_callout(
        doc,
        "What “merged AI” looks like in a working week",
        "A district manager opens a physician already stored in GNP. Agentforce returns a "
        "Visit Insights brief (Cloudastick skill) and any GNP-native score GNP already "
        "computes. If sell-out on an affiliated pharmacy moved, the Planning skill proposes "
        "three follow-ups. The manager accepts two. Draft visits appear in GNP’s tracker. "
        "One conversation, two agent families, one system of record.",
    )


def add_raci(doc: Document) -> None:
    doc.add_heading("7.  What stays with GNP, what we add", level=1)
    add_table(
        doc,
        ["Capability", "Stays in GNP", "Added by Agentforce + Cloudastick"],
        [
            (
                "Visit capture and tracking",
                "Yes — existing process and screens",
                "Visit Insights briefs, pace, nearby follow-up, apply-to-plan",
            ),
            (
                "Physician reporting",
                "Yes — existing reports",
                "Agentforce interpretation, APM, next-best action from the same facts",
            ),
            (
                "Physician / pharmacy / HCO master",
                "GNP holds operational copies",
                "Live Cloudastick master, Search Before Create, Map Intelligence",
            ),
            (
                "GNP AI module",
                "Remains; becomes a peer skill",
                "Agentforce Life Sciences fabric, Trust Layer, Cloudastick pharma skills",
            ),
            (
                "Distributor / sell-out files",
                "GNP continues to receive them",
                "Product-code unification, pharmacy matching, quality-gated load",
            ),
            (
                "Field-force identity and org structure",
                "GNP remains source unless agreed otherwise",
                "Mapped into Data 360 so agents know territory and role",
            ),
        ],
        col_widths=[2.0, 2.2, 2.5],
    )


def add_delivery(doc: Document) -> None:
    doc.add_heading("8.  Indicative delivery approach", level=1)
    add_body(
        doc,
        "Durations below are orientation only. They assume GNP can nominate integration "
        "owners and provide sample visit, account, and distributor extracts. A discovery "
        "workshop will replace these ranges with a dated plan. No commercial figures are "
        "included in this document on purpose.",
    )
    add_table(
        doc,
        ["Phase", "Focus", "GNP sees", "Depends on"],
        [
            (
                "0  Discovery",
                "Interfaces, identity, AI inventory, data samples, compliance",
                "Agreed architecture, skill backlog, integration map",
                "GNP IT and commercial sponsors in the room",
            ),
            (
                "1  Master",
                "Search Before Create on physician, pharmacy, HCO",
                "Create flow that searches Cloudastick before insert; fewer duplicates",
                "Create-API or embed point in GNP; approval rule if GNP uses one",
            ),
            (
                "2  Visit Insights",
                "Agentforce on existing tracking and physician reports",
                "Pre-call briefs, pace, nearby follow-up — inside GNP or a sidecar",
                "Read access to visit and account history",
            ),
            (
                "3  Market + APM",
                "Distributor unification and Absolute Physician Movement",
                "One product language; movement insights on the physicians GNP already rates",
                "Sample wholesaler files and product catalogue mapping",
            ),
            (
                "4  Agent fabric",
                "Register GNP AI as a peer skill; apply-actions under GNP approvals",
                "One suggestion list from both AI modules, audited in Agentforce",
                "Phases 1–2 stable; GNP AI contract documented",
            ),
        ],
        col_widths=[1.25, 1.7, 1.9, 1.85],
    )
    add_body(
        doc,
        "Each phase is designed to leave GNP better off even if the next phase waits. "
        "Search Before Create has value without APM. Visit Insights has value without "
        "distributor files. Nothing in phase 4 requires GNP to retire its own AI.",
        size=10.5,
        color=GRAY,
        italic=True,
    )


def add_trust(doc: Document) -> None:
    doc.add_heading("9.  Trust, compliance, and data use", level=1)
    add_body(
        doc,
        "Life sciences AI is only useful if medical, legal, and IT can explain it. "
        "Agentforce Life Sciences is proposed because the Einstein Trust Layer is already "
        "the Salesforce control plane for that conversation — not because it is fashionable.",
    )
    add_bullet(doc, " Prompts and responses can be audited. Grounding is on GNP and master data, not an open browse of the web.")
    add_bullet(doc, " GNP confidential data is not used to train foundation models.")
    add_bullet(doc, " Public professional retrieval is source-limited and can be switched off per skill.")
    add_bullet(doc, " Apply-actions respect GNP approvals, territory visibility, and existing field-force roles.")
    add_bullet(doc, " Master identity writes back only the field set GNP authorises.")
    add_bullet(doc, " Hosting, residency, and sandbox strategy are confirmed in discovery against GNP policy.")
    add_body(
        doc,
        "Cloudastick will not claim GxP validation of Agentforce itself in this proposal. "
        "If GNP requires a validated path for a specific process, that is scoped as its own "
        "workstream after discovery — it is not assumed.",
        before=6,
        italic=True,
        size=10.5,
        color=GRAY,
    )


def add_next(doc: Document) -> None:
    doc.add_heading("10.  Suggested next step", level=1)
    add_body(
        doc,
        "A half-day working session is enough to turn this document into a buildable plan. "
        "The session should include GNP commercial operations (or SFE), the owner of the "
        "in-house platform, and whoever owns GNP’s current AI module.",
    )
    add_table(
        doc,
        ["Workshop item", "Why it matters"],
        [
            ("Walk the current create-physician / create-pharmacy path", "Defines where Search Before Create is invoked."),
            ("Sample of visit and physician-report extracts (anonymised if needed)", "Proves Visit Insights can run on GNP data as-is."),
            ("Inventory of GNP’s AI outputs and APIs", "Lets us register GNP AI as a peer Agentforce skill."),
            ("One distributor file and the product-code list behind it", "Sizes unification effort honestly."),
            ("Decision: embed in GNP UI, Salesforce sidecar, or both for stewards", "Keeps the “do not rebuild” rule intact."),
        ],
        col_widths=[3.2, 3.5],
    )
    add_callout(
        doc,
        "Close",
        "GNP has already done the hard part: the field captures visits and reports on "
        "physicians. Cloudastick’s proposal is to put Salesforce Agentforce Life Sciences "
        "and a live customer master underneath that work, and to let GNP’s own AI remain "
        "part of the same fabric. Add intelligence. Do not start again.",
    )
    add_body(
        doc,
        "Cloudastick Systems  ·  Salesforce Partner",
        bold=True,
        color=NAVY,
        after=2,
    )
    add_body(
        doc,
        f"Document date: {date.today():%d %B %Y}. This Word file is intended to be edited. "
        "Add GNP system names, tighten wording for the audience, and attach commercials "
        "separately if required.",
        size=10,
        color=MUTED,
        italic=True,
    )


def add_appendix(doc: Document) -> None:
    doc.add_heading("Appendix  —  Agent catalogue", level=1)
    add_body(
        doc,
        "The catalogue below is how Cloudastick pharmaceutical intelligence is branded and "
        "governed under Agentforce Life Sciences. Names can be localised for GNP. Skills "
        "are independently switchable; GNP does not have to take the full set.",
    )
    add_table(
        doc,
        ["Agentforce skill", "Business question it answers", "Typical apply-action"],
        [
            (
                "Search Before Create",
                "Does this physician, pharmacy, or HCO already exist in the live master?",
                "Adopt master identity or open a quality-scored create",
            ),
            (
                "Map Intelligence",
                "If the master has no row, is there a locatable professional site we should seed from?",
                "Pre-fill create; never silent insert",
            ),
            (
                "Physician Intelligence",
                "Who is this HCP, and do our ratings still match public professional context and visits?",
                "Flag mismatch; steward or rep confirms",
            ),
            (
                "Visit Insights",
                "What should I know before this call, and am I on pace?",
                "Draft nearby visits; open prior call",
            ),
            (
                "Absolute Physician Movement",
                "Did this physician actually move on the brand, or only in the rating?",
                "Coach, retarget, or correct classification",
            ),
            (
                "Market Unification",
                "Are these distributor codes the same product and the same outlet?",
                "Confirm match; load sell-out onto the right account",
            ),
            (
                "Planning / Next Best Action",
                "Given sell-out and affiliations, whom should we visit in this cycle?",
                "Create accepted draft visits in GNP",
            ),
            (
                "Account & Home Briefs",
                "What is the territory gap today, and what is the story on this account?",
                "Surface in GNP or Salesforce; optional plan update",
            ),
        ],
        col_widths=[1.85, 2.7, 2.15],
    )
    add_body(
        doc,
        "Optional later skills (not assumed in phases 1–2): closed-loop marketing insights, "
        "sample-compliance prompts, coaching-gap briefs, and executive roll-ups. They follow "
        "the same pattern — Agentforce reads GNP data, suggests, and applies only with permission.",
        size=10.5,
        color=GRAY,
    )


def set_core_properties(doc: Document) -> None:
    props = doc.core_properties
    props.author = "Cloudastick Systems"
    props.last_modified_by = "Cloudastick Systems"
    props.title = "Agentforce Life Sciences for Global Napi Pharmaceuticals"
    props.subject = "Salesforce Agentforce Life Sciences + Cloudastick intelligence layer for GNP"
    props.category = "Proposal"
    props.comments = (
        "Discussion document. Editable. Does not replace GNP’s in-house platform. "
        "No commercial figures included."
    )


def build_document() -> Document:
    doc = Document()
    configure_styles(doc)
    set_core_properties(doc)

    section = doc.sections[0]
    configure_section(section, different_first=True)
    add_header_footer(section, first_page=True)
    add_header_footer(section, first_page=False)

    add_cover(doc)
    add_contents(doc)
    add_heard(doc)
    add_principle(doc)
    add_architecture(doc)
    add_search_before_create(doc)
    add_agentforce_layer(doc)
    add_merge(doc)
    add_raci(doc)
    add_delivery(doc)
    add_trust(doc)
    add_next(doc)
    add_appendix(doc)
    return doc


def main() -> None:
    doc = build_document()
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
